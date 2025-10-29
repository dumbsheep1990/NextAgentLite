"""
DOCX文件解析器
"""

from typing import List
from .base_parser import BaseParser
from models.document_chunk import ContentType, DocumentBlock


class DocxParser(BaseParser):
    """DOCX文件解析器"""
    
    def __init__(self):
        super().__init__()
        self.supported_formats = ['.docx', '.doc']

    def parse(self, file_path: str, **kwargs) -> List[DocumentBlock]:
        """
        解析DOCX文件

        Args:
            file_path: 文件路径
            **kwargs: 解析参数

        Returns:
            文档块列表
        """
        import os

        # 检查文件扩展名
        file_ext = os.path.splitext(file_path)[1].lower()

        if file_ext == '.doc':
            # .doc 是老格式，python-docx不支持，使用回退方案
            return self._parse_doc_file(file_path)

        try:
            # .docx 使用python-docx
            from docx import Document
            return self._parse_with_python_docx(file_path)
        except ImportError:
            # 如果没有安装python-docx，返回占位符
            return self._create_placeholder_block(file_path)
        except Exception as e:
            # 解析失败，使用回退方案
            import logging
            logging.warning(f"DocxParser解析失败: {e}，使用回退方案")
            return self._parse_doc_file(file_path)
    
    def _parse_with_python_docx(self, file_path: str) -> List[DocumentBlock]:
        """使用python-docx解析DOCX"""
        from docx import Document
        from docx.shared import Inches
        
        doc = Document(file_path)
        blocks = []
        current_headings = []
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            
            if not text:
                continue
            
            # 检测标题样式
            style_name = paragraph.style.name.lower()
            
            if 'heading' in style_name:
                # 提取标题级别
                try:
                    level = int(style_name.split()[-1]) if style_name.split()[-1].isdigit() else 1
                except:
                    level = 1
                
                # 更新标题堆栈
                current_headings = current_headings[:level-1] + [text]
                
                # 标题也作为内容块
                block = self.create_document_block(
                    content=text,
                    content_type=ContentType.TEXT,
                    headings=current_headings.copy(),
                    extra={'style': style_name, 'is_heading': True, 'level': level}
                )
                blocks.append(block)
                
            else:
                # 普通段落
                processed_text = self.preprocess_content(text)
                
                block = self.create_document_block(
                    content=processed_text,
                    content_type=ContentType.TEXT,
                    headings=current_headings.copy(),
                    extra={'style': style_name}
                )
                blocks.append(block)
        
        # 处理表格
        for table in doc.tables:
            table_text = self._extract_table_text(table)
            if table_text.strip():
                block = self.create_document_block(
                    content=table_text,
                    content_type=ContentType.TABLE,
                    headings=current_headings.copy(),
                    extra={'type': 'table'}
                )
                blocks.append(block)
        
        return blocks
    
    def _extract_table_text(self, table) -> str:
        """提取表格文本"""
        table_data = []
        
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                row_data.append(cell_text)
            table_data.append(" | ".join(row_data))
        
        return "\n".join(table_data)
    
    def _parse_doc_file(self, file_path: str) -> List[DocumentBlock]:
        """
        解析 .doc 文件（Word 97-2003 格式）
        使用多种回退方案
        """
        import os
        import logging

        logger = logging.getLogger(__name__)

        # 方案1: 尝试使用 textract（如果已安装）
        try:
            import textract
            text = textract.process(file_path).decode('utf-8')
            logger.info(f"使用 textract 成功解析 .doc 文件: {os.path.basename(file_path)}")
            return self._create_text_blocks(text)
        except ImportError:
            logger.debug("textract 未安装，尝试其他方案")
        except Exception as e:
            logger.warning(f"textract 解析失败: {e}")

        # 方案2: 尝试使用 antiword（Linux/Mac 命令行工具）
        try:
            import subprocess
            result = subprocess.run(
                ['antiword', file_path],
                capture_output=True,
                text=True,
                timeout=30
            )
            if result.returncode == 0 and result.stdout:
                text = result.stdout
                logger.info(f"使用 antiword 成功解析 .doc 文件: {os.path.basename(file_path)}")
                return self._create_text_blocks(text)
        except FileNotFoundError:
            logger.debug("antiword 未安装")
        except Exception as e:
            logger.warning(f"antiword 解析失败: {e}")

        # 方案3: 简单的二进制文本提取（最后的回退）
        logger.warning(f".doc 文件解析降级为简单文本提取: {os.path.basename(file_path)}")
        try:
            with open(file_path, 'rb') as f:
                binary_data = f.read()

            # 尝试提取可打印的文本
            text = binary_data.decode('latin-1', errors='ignore')
            # 过滤非打印字符，保留常见标点
            import string
            printable = set(string.printable)
            text = ''.join(c for c in text if c in printable or ord(c) > 127)

            # 清理连续空白
            import re
            text = re.sub(r'\s+', ' ', text)
            text = text.strip()

            if len(text) < 50:
                # 提取的文本太少，返回错误提示
                return self._create_error_block(file_path, ".doc 格式解析失败，建议转换为 .docx 格式")

            logger.info(f"简单文本提取完成: {os.path.basename(file_path)}, 长度: {len(text)}")
            return self._create_text_blocks(text)

        except Exception as e:
            logger.error(f"简单文本提取也失败: {e}")
            return self._create_error_block(file_path, f".doc 文件解析失败: {str(e)}")

    def _create_text_blocks(self, text: str) -> List[DocumentBlock]:
        """从文本创建文档块"""
        if not text or not text.strip():
            return []

        # 按段落分割（保留一定的段落结构）
        paragraphs = [p.strip() for p in text.split('\n') if p.strip()]

        blocks = []
        for para in paragraphs:
            if len(para) < 10:  # 跳过太短的段落
                continue

            block = self.create_document_block(
                content=self.preprocess_content(para),
                content_type=ContentType.TEXT,
                extra={'source': 'doc_fallback'}
            )
            blocks.append(block)

        return blocks if blocks else [self._create_error_block("unknown", "无法提取有效文本")]

    def _create_error_block(self, file_path: str, error_message: str) -> List[DocumentBlock]:
        """创建错误提示块"""
        import os
        filename = os.path.basename(file_path)

        block = self.create_document_block(
            content=f"文件: {filename}\n错误: {error_message}\n\n建议：请将 .doc 文件转换为 .docx 格式后重新上传，以获得最佳解析效果。",
            content_type=ContentType.TEXT,
            extra={'source': 'error', 'error': error_message}
        )

        return [block]

    def _create_placeholder_block(self, file_path: str) -> List[DocumentBlock]:
        """创建占位符块（当没有python-docx库时）"""
        import os
        filename = os.path.basename(file_path)

        block = self.create_document_block(
            content=f"DOCX文件: {filename}\n注意：需要安装python-docx库来解析DOCX内容",
            content_type=ContentType.TEXT,
            page_number=1,
            extra={'source': 'placeholder', 'error': 'no_docx_library'}
        )

        return [block]